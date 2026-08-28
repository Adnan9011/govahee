from __future__ import annotations

import csv
import io
from datetime import date

from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from certificates.variables import BUILT_IN_VARIABLES

ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".xls"}
MAX_ROWS = 10000


def parse_tabular(file_obj, filename: str) -> tuple[list[str], list[dict]]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return _parse_csv(file_obj)
    if name.endswith(".xlsx") or name.endswith(".xls"):
        return _parse_xlsx(file_obj)
    raise ValidationError("فقط فایل CSV یا Excel مجاز است.")


def _parse_csv(file_obj) -> tuple[list[str], list[dict]]:
    raw = file_obj.read()
    if isinstance(raw, bytes):
        text = raw.decode("utf-8-sig", errors="replace")
    else:
        text = raw
    reader = csv.DictReader(io.StringIO(text))
    headers = [h.strip() for h in (reader.fieldnames or [])]
    rows = []
    for i, row in enumerate(reader, start=2):
        rows.append({"row_number": i, "data": {k.strip(): (v or "").strip() for k, v in row.items() if k}})
        if len(rows) > MAX_ROWS:
            raise ValidationError("تعداد ردیف‌ها از حد مجاز بیشتر است.")
    return headers, rows


def _parse_xlsx(file_obj) -> tuple[list[str], list[dict]]:
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return [], []
    headers = [str(c).strip() if c is not None else "" for c in header_row]
    result = []
    for i, row in enumerate(rows_iter, start=2):
        data = {}
        empty = True
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = row[idx] if idx < len(row) else ""
            if value is None:
                value = ""
            elif isinstance(value, date):
                value = value.isoformat()
            else:
                value = str(value).strip()
            data[header] = value
            if value:
                empty = False
        if empty:
            continue
        result.append({"row_number": i, "data": data})
        if len(result) > MAX_ROWS:
            raise ValidationError("تعداد ردیف‌ها از حد مجاز بیشتر است.")
    return [h for h in headers if h], result


def apply_mapping(payload: dict, mapping: dict[str, str]) -> dict[str, str]:
    mapped: dict[str, str] = {}
    for col, field in (mapping or {}).items():
        if not field:
            continue
        mapped[field] = str((payload or {}).get(col, "") or "").strip()
    return mapped


def validate_mapped_row(mapped: dict[str, str], *, row_number: int) -> list[str]:
    errors: list[str] = []
    name = mapped.get("recipient_name") or mapped.get("name") or ""
    if not name:
        errors.append(f"Row {row_number}: نام دریافت‌کننده الزامی است.")
    email = mapped.get("email") or ""
    if email and "@" not in email:
        errors.append(f"Row {row_number}: Email is invalid")
    issue_date = mapped.get("issue_date") or ""
    if issue_date:
        try:
            date.fromisoformat(issue_date.replace("/", "-")[:10])
        except ValueError:
            errors.append(f"Row {row_number}: تاریخ صدور نامعتبر است.")
    return errors


def validate_batch_items(items, mapping: dict[str, str]) -> list[dict]:
    report = []
    for item in items:
        mapped = apply_mapping(item.payload, mapping)
        errors = validate_mapped_row(mapped, row_number=item.row_number)
        item.errors = errors
        item.status = (
            item.Status.INVALID if errors else item.Status.VALID
        )
        if errors:
            report.append({"row": item.row_number, "errors": errors})
    return report


def suggested_mapping(headers: list[str]) -> dict[str, str]:
    aliases = {
        "name": "recipient_name",
        "full_name": "recipient_name",
        "recipient": "recipient_name",
        "نام": "recipient_name",
        "email": "email",
        "ایمیل": "email",
        "course": "course_name",
        "course_name": "course_name",
        "دوره": "course_name",
        "score": "score",
        "نمره": "score",
        "date": "issue_date",
        "issue_date": "issue_date",
        "تاریخ": "issue_date",
    }
    mapping = {}
    for header in headers:
        key = header.strip().lower()
        if key in aliases:
            mapping[header] = aliases[key]
        elif header in BUILT_IN_VARIABLES:
            mapping[header] = header
    return mapping


def sanitize_csv_value(value: str) -> str:
    if not value:
        return value
    if value[0] in ("=", "+", "-", "@"):
        return "'" + value
    return value
